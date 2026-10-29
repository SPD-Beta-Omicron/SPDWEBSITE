import io
from datetime import date
from unittest import expectedFailure

from django.urls import reverse
from openpyxl import load_workbook

from core.testing import SPDTestCase
from .models import Philanthropy_Hours_Event_and_Request as Event
from .templatetags.philanthropy_extras import hours_remaining
from .views_functions import get_academic_year_dates


class PhilanthropyTestCase(SPDTestCase):

    def setUp(self):
        self.active = self.make_member('active', 'ACTIVE', first='Alex', last='Active')
        self.chair = self.make_member('chair', 'ACTIVE', 'PHIL_CHAIR', first='Casey', last='Chair')
        self.alumnus = self.make_member('alumnus', 'ALUMNI')

    def event(self, user, hours, when, status='approved', title='Food bank'):
        return Event.objects.create(user=user, philanthropy_event_title=title, philanthropy_event_hours=hours,
                                    philanthropy_event_date=when, philanthropy_approval_status=status)


class AccessTests(PhilanthropyTestCase):

    def test_logged_out_users_are_sent_to_login(self):
        for url in [reverse('philanthropy_dashboard'), reverse('philanthropy_request'),
                    reverse('philanthropy_approve'), reverse('export_philanthropy_hours')]:
            with self.subTest(url=url):
                self.assertRedirectsToLogin(self.client.get(url))

    def test_alumni_cannot_see_dashboard(self):
        self.login(self.alumnus)
        self.assertDenied(self.client.get(reverse('philanthropy_dashboard')))

    def test_only_chair_can_approve_or_edit(self):
        event = self.event(self.active, 3, date(2025, 9, 1), status='requested')
        self.login(self.active)
        self.assertDenied(self.client.get(reverse('philanthropy_approve')))
        self.assertDenied(self.client.get(reverse('philanthropy_edit_event', args=[event.id])))
        self.login(self.chair)
        self.assertEqual(self.client.get(reverse('philanthropy_approve')).status_code, 200)
        self.assertEqual(self.client.get(reverse('philanthropy_edit_event', args=[event.id])).status_code, 200)

    @expectedFailure
    def test_only_chair_can_export_hours(self):
        """KNOWN BUG: the role check on the export view is commented out, so any logged-in user can download everyone's hours."""
        self.login(self.active)
        self.assertDenied(self.client.get(reverse('export_philanthropy_hours')))


class DashboardTests(PhilanthropyTestCase):

    def test_totals_for_selected_academic_year(self):
        self.event(self.active, 6, date(2025, 9, 1))
        self.event(self.active, 5, date(2026, 3, 1))
        self.event(self.chair, 4, date(2025, 10, 1))
        self.event(self.chair, 20, date(2025, 10, 2), status='requested')   # not approved yet
        self.event(self.chair, 20, date(2024, 10, 1))                       # previous year
        self.login(self.active)
        context = self.client.get(reverse('philanthropy_dashboard'), {'year': 2025}).context

        self.assertEqual(context['total_chapter_hours'], 15)
        self.assertEqual(context['members_meeting_goal'], 1)      # only Alex has 10+
        self.assertEqual(context['total_active_members'], 2)
        self.assertEqual(context['avg_hours_per_brother'], 7.5)
        self.assertEqual([b['full_name'] for b in context['brothers_progress']], ['Alex Active', 'Casey Chair'])
        self.assertEqual(context['user_total_hours'], 11)
        self.assertEqual(context['year_label'], '2025–2026 Academic Year')

    def test_chair_flag(self):
        self.login(self.active)
        self.assertFalse(self.client.get(reverse('philanthropy_dashboard')).context['is_philanthropy_chair'])
        self.login(self.chair)
        self.assertTrue(self.client.get(reverse('philanthropy_dashboard')).context['is_philanthropy_chair'])

    def test_invalid_year_does_not_crash(self):
        """A mistyped ?year= should not take the page down."""
        self.login(self.active)
        self.client.raise_request_exception = False
        response = self.client.get(reverse('philanthropy_dashboard'), {'year': 'abc'})
        self.assertNotEqual(response.status_code, 500)

    test_invalid_year_does_not_crash = expectedFailure(test_invalid_year_does_not_crash)
    test_invalid_year_does_not_crash.__doc__ = "KNOWN BUG: ?year=abc raises ValueError and returns a server error."


class RequestTests(PhilanthropyTestCase):

    def submit(self, hours, title='Park cleanup', when='2025-09-15'):
        self.login(self.active)
        return self.client.post(reverse('philanthropy_request'),
                                {'event_title': title, 'hours': hours, 'event_date': when})

    def test_request_form_loads(self):
        self.login(self.active)
        self.assertEqual(self.client.get(reverse('philanthropy_request')).status_code, 200)

    def test_submitting_hours_creates_pending_request(self):
        response = self.submit('3')
        self.assertRedirects(response, reverse('philanthropy_dashboard'), fetch_redirect_response=False)
        event = Event.objects.get()
        self.assertEqual((event.user, event.philanthropy_event_hours, event.philanthropy_approval_status),
                         (self.active, 3, 'requested'))

    def test_less_than_half_an_hour_is_rejected(self):
        response = self.submit('0.25')
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Event.objects.exists())
        self.assertTrue(any('at least 0.5 hours' in m for m in self.messages_for(response)))

    @expectedFailure
    def test_half_hours_are_kept(self):
        """KNOWN BUG: hours are stored in an IntegerField, so 1.5 hours is saved as 1 (and 0.5 as 0)."""
        self.submit('1.5')
        self.assertEqual(Event.objects.get().philanthropy_event_hours, 1.5)


class ApprovalTests(PhilanthropyTestCase):

    def setUp(self):
        super().setUp()
        self.pending = self.event(self.active, 3, date(2025, 9, 1), status='requested')
        self.login(self.chair)

    def test_pending_list(self):
        response = self.client.get(reverse('philanthropy_approve'))
        self.assertEqual(list(response.context['pending_requests']), [self.pending])

    def test_approve(self):
        self.client.post(reverse('philanthropy_approve'), {'event_id': self.pending.id, 'action': 'approve'})
        self.pending.refresh_from_db()
        self.assertEqual(self.pending.philanthropy_approval_status, 'approved')
        self.assertEqual(self.pending.philanthropy_event_approver_name, 'Casey Chair')
        self.assertIsNotNone(self.pending.philanthropy_event_submission_approval_date)

    def test_deny(self):
        self.client.post(reverse('philanthropy_approve'), {'event_id': self.pending.id, 'action': 'deny'})
        self.pending.refresh_from_db()
        self.assertEqual(self.pending.philanthropy_approval_status, 'denied')

    def test_unknown_event(self):
        response = self.client.post(reverse('philanthropy_approve'), {'event_id': 99999, 'action': 'approve'})
        self.assertIn('Event not found', self.messages_for(response))

    def test_edit_event(self):
        response = self.client.post(reverse('philanthropy_edit_event', args=[self.pending.id]),
                                    {'event_title': 'Renamed', 'hours': '4', 'event_date': '2025-09-02'})
        self.assertRedirects(response, reverse('brother_philanthropy_history', args=[self.active.id]),
                             fetch_redirect_response=False)
        self.pending.refresh_from_db()
        self.assertEqual((self.pending.philanthropy_event_title, self.pending.philanthropy_event_hours), ('Renamed', 4))
        self.assertEqual(self.pending.last_edited_by, self.chair)


class HistoryAndExportTests(PhilanthropyTestCase):

    def test_history_grouped_by_academic_year(self):
        self.event(self.active, 2, date(2025, 9, 1))
        self.event(self.active, 3, date(2026, 6, 30))     # still the 2025–26 year
        self.event(self.active, 4, date(2024, 8, 1))
        self.event(self.active, 9, date(2025, 9, 2), status='denied')
        self.login(self.active)
        context = self.client.get(reverse('brother_philanthropy_history', args=[self.active.id])).context
        years = [(y['start'], y['total_hours']) for y in context['academic_year_totals']]
        self.assertEqual(years, [(2025, 5), (2024, 4)])
        self.assertEqual(context['lifetime_total'], 9)

    def test_history_for_unknown_brother_redirects(self):
        self.login(self.active)
        response = self.client.get(reverse('brother_philanthropy_history', args=[99999]))
        self.assertRedirects(response, reverse('philanthropy_dashboard'), fetch_redirect_response=False)

    def test_export_spreadsheet(self):
        self.event(self.active, 6, date(2025, 9, 1))
        self.event(self.active, 5, date(2025, 9, 2))
        self.event(self.chair, 2, date(2025, 9, 3), status='requested')
        self.login(self.chair)
        response = self.client.get(reverse('export_philanthropy_hours'))
        self.assertIn('approved_philanthropy_hours.xlsx', response['Content-Disposition'])
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(row=1, column=1).value, 'Alex Active')
        self.assertEqual(sheet.cell(row=2, column=1).value, 'Total: 11')
        self.assertEqual(sheet.cell(row=1, column=6).value, None)    # pending hours are not exported


class HelperTests(SPDTestCase):

    def test_academic_year_runs_july_to_june(self):
        self.assertEqual(get_academic_year_dates(2025), (date(2025, 7, 1), date(2026, 6, 30)))

    def test_hours_remaining_filter(self):
        self.assertEqual(hours_remaining(10, 4), 6)
        self.assertEqual(hours_remaining(10, 12), 0)
