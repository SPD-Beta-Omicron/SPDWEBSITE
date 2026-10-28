import io
from datetime import date, timedelta

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from CORE.testing import SPDTestCase
from .models import NewMember_Mark_Event_and_Request as Mark
from .views_functions import calculate_member_marks, get_new_members, get_user_submitted_marks


class NewMemberTestCase(SPDTestCase):

    def setUp(self):
        self.active = self.make_member('active', 'ACTIVE', first='Alex', last='Active')
        self.board = self.make_member('board', 'ACTIVE', 'NM_BOARD', first='Blair', last='Board')
        self.educator = self.make_member('educator', 'ACTIVE', 'NM_BOARD', 'NME', first='Eden', last='Educator')
        self.nm1 = self.make_member('nm1', 'NEW_MEMBER', first='Noah', last='Able', pclass='Current')
        self.nm2 = self.make_member('nm2', 'NEW_MEMBER', first='Nina', last='Baker', pclass='Current')

    def mark(self, target, value, status='approved', when=None, requester=None, title='Helped clean'):
        return Mark.objects.create(requesting_user=requester or self.active, target_user=target,
                                   mark_event_title=title, mark_value=value,
                                   mark_event_date=when or date(2025, 9, 1), mark_approval_status=status)


class AccessTests(NewMemberTestCase):

    def test_logged_out_users_are_sent_to_login(self):
        for url in [reverse('newmember_marks_dashboard'), reverse('newmember_request'), reverse('newmember_approve'),
                    reverse('export_newmember_marks')]:
            with self.subTest(url=url):
                self.assertRedirectsToLogin(self.client.get(url))

    def test_new_member_landing_page_loads(self):
        self.login(self.nm1)
        self.assertEqual(self.client.get(reverse('newmember_dashboard')).status_code, 200)

    def test_new_members_cannot_submit_marks(self):
        self.login(self.nm1)
        response = self.client.get(reverse('newmember_request'))
        self.assertRedirects(response, reverse('newmember_marks_dashboard'), fetch_redirect_response=False)
        self.assertIn('Only active members can submit marks.', self.messages_for(response))

    def test_only_board_can_approve_or_edit(self):
        mark = self.mark(self.nm1, 1, status='requested')
        self.login(self.active)
        self.assertEqual(self.client.get(reverse('newmember_approve')).status_code, 302)
        self.assertEqual(self.client.get(reverse('newmember_edit_mark', args=[mark.id])).status_code, 302)
        self.login(self.board)
        self.assertEqual(self.client.get(reverse('newmember_approve')).status_code, 200)
        self.assertEqual(self.client.get(reverse('newmember_edit_mark', args=[mark.id])).status_code, 200)


class DashboardTests(NewMemberTestCase):

    def test_active_member_sees_chapter_analytics(self):
        self.mark(self.nm1, 3)
        self.mark(self.nm2, -2)
        self.mark(self.nm2, 1, status='requested')
        self.login(self.active)
        response = self.client.get(reverse('newmember_marks_dashboard'))
        self.assertTemplateUsed(response, 'NEWMEMBER/ACTIVE_newmember_marks_dashboard.html')
        analytics = response.context['analytics']
        self.assertEqual(analytics['total_new_members'], 2)
        self.assertEqual(analytics['members_in_good_standing'], 1)
        self.assertEqual(analytics['average_marks'], 0.5)
        self.assertEqual(analytics['pending_requests'], 1)
        self.assertIsNone(response.context['personal_stats'])

    def test_new_member_sees_personal_stats(self):
        self.mark(self.nm1, 4)
        self.mark(self.nm2, 2)
        self.login(self.nm1)
        response = self.client.get(reverse('newmember_marks_dashboard'))
        self.assertTemplateUsed(response, 'NEWMEMBER/newmember_marks_dashboard.html')
        stats = response.context['personal_stats']
        self.assertEqual((stats['total_marks'], stats['class_average'], stats['standing']), (4, 3.0, 'Good Standing'))

    def test_board_flag(self):
        self.login(self.board)
        self.assertTrue(self.client.get(reverse('newmember_marks_dashboard')).context['is_new_member_board'])


class SubmitAndApproveTests(NewMemberTestCase):

    def submit(self, user, targets, value=2):
        self.login(user)
        return self.client.post(reverse('newmember_request'), {
            'target_users': [t.id for t in targets], 'mark_value': value,
            'mark_reason': 'Showed up early', 'mark_date': '2025-09-10'})

    def test_request_form_lists_new_members(self):
        self.login(self.active)
        response = self.client.get(reverse('newmember_request'))
        self.assertEqual([m.user for m in response.context['new_members']], [self.nm1, self.nm2])

    def test_active_member_submission_needs_approval(self):
        response = self.submit(self.active, [self.nm1, self.nm2])
        self.assertRedirects(response, reverse('newmember_marks_dashboard'), fetch_redirect_response=False)
        self.assertEqual(Mark.objects.filter(mark_approval_status='requested').count(), 2)

    def test_board_submission_is_approved_immediately(self):
        self.submit(self.board, [self.nm1])
        mark = Mark.objects.get()
        self.assertEqual((mark.mark_approval_status, mark.mark_approver_name), ('approved', 'Blair Board'))

    def test_board_approves_and_denies(self):
        first = self.mark(self.nm1, 1, status='requested')
        second = self.mark(self.nm2, 1, status='requested')
        self.login(self.board)
        self.client.post(reverse('newmember_approve'), {'mark_id': first.id, 'action': 'approve'})
        self.client.post(reverse('newmember_approve'), {'mark_id': second.id, 'action': 'deny'})
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual((first.mark_approval_status, first.mark_approver_name), ('approved', 'Blair Board'))
        self.assertEqual(second.mark_approval_status, 'denied')

    def test_approving_unknown_mark(self):
        self.login(self.board)
        response = self.client.post(reverse('newmember_approve'), {'mark_id': 99999, 'action': 'approve'})
        self.assertIn('Mark not found', self.messages_for(response))

    def test_submitted_marks_hide_approved_and_old_denied(self):
        recent = timezone.now().date() - timedelta(days=3)
        old = timezone.now().date() - timedelta(days=30)
        pending = self.mark(self.nm1, 1, status='requested', when=recent)
        recent_denied = self.mark(self.nm1, 1, status='denied', when=recent)
        self.mark(self.nm1, 1, status='denied', when=old)
        self.mark(self.nm1, 1, status='approved', when=recent)
        self.assertEqual(set(get_user_submitted_marks(self.active)), {pending, recent_denied})


class EditTests(NewMemberTestCase):

    def setUp(self):
        super().setUp()
        self.existing = self.mark(self.nm1, 1)

    def test_board_edits_mark(self):
        self.login(self.board)
        response = self.client.post(reverse('newmember_edit_mark', args=[self.existing.id]), {
            'mark_reason': 'Corrected', 'mark_value': -1, 'mark_date': '2025-09-02', 'target_user': self.nm2.id})
        self.assertRedirects(response, reverse('newmember_mark_history', args=[self.nm2.id]), fetch_redirect_response=False)
        self.existing.refresh_from_db()
        self.assertEqual((self.existing.mark_event_title, self.existing.mark_value, self.existing.target_user),
                         ('Corrected', -1, self.nm2))
        self.assertEqual(self.existing.last_edited_by, self.board)

    def test_only_educator_can_deny_from_edit_page(self):
        self.login(self.board)
        self.client.post(reverse('newmember_edit_mark', args=[self.existing.id]), {'action': 'deny'})
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.mark_approval_status, 'approved')

        self.login(self.educator)
        self.client.post(reverse('newmember_edit_mark', args=[self.existing.id]), {'action': 'deny'})
        self.existing.refresh_from_db()
        self.assertEqual((self.existing.mark_approval_status, self.existing.mark_approver_name), ('denied', 'Eden Educator'))


class HistoryTests(NewMemberTestCase):

    def test_history_totals_only_approved_marks(self):
        self.mark(self.nm1, -3)
        self.mark(self.nm1, 1)
        self.mark(self.nm1, 10, status='requested')
        self.login(self.active)
        context = self.client.get(reverse('newmember_mark_history', args=[self.nm1.id])).context
        self.assertEqual((context['total_marks'], context['is_negative']), (-2, True))
        self.assertEqual(len(context['marks']), 2)

    def test_history_for_someone_who_is_not_a_new_member(self):
        self.login(self.active)
        response = self.client.get(reverse('newmember_mark_history', args=[self.active.id]))
        self.assertRedirects(response, reverse('newmember_dashboard'), fetch_redirect_response=False)

    def test_progress_is_sorted_lowest_first(self):
        self.mark(self.nm1, 5)
        self.mark(self.nm2, -1)
        progress = calculate_member_marks(get_new_members())
        self.assertEqual([(p['user'], p['total_marks']) for p in progress], [(self.nm2, -1), (self.nm1, 5)])


class ExportTests(NewMemberTestCase):

    def export(self, user, pclass):
        self.login(user)
        return self.client.get(reverse('export_newmember_marks'), {'pclass': pclass})

    def test_export_current_class(self):
        self.mark(self.nm1, 2)
        self.mark(self.nm1, 3)
        self.mark(self.nm2, 9, status='requested')
        response = self.export(self.active, 'Current')
        self.assertIn('Approved_Current_Class_Newmember_Marks.xlsx', response['Content-Disposition'])
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(row=1, column=1).value, 'Noah Able')
        self.assertEqual(sheet.cell(row=2, column=1).value, 'Total: 5')
        self.assertIsNone(sheet.cell(row=1, column=6).value)    # Nina has no approved marks

    def test_export_past_class_by_pclass(self):
        alum = self.make_member('alum', 'ACTIVE', first='Old', last='Timer', pclass='Beta Psi')
        self.mark(alum, 4)
        self.mark(self.nm1, 2)
        sheet = load_workbook(io.BytesIO(self.export(self.active, 'Beta Psi').content)).active
        self.assertEqual(sheet.cell(row=1, column=1).value, 'Old Timer')
        self.assertIsNone(sheet.cell(row=1, column=6).value)

    def test_new_members_can_only_export_current_class(self):
        response = self.export(self.nm1, 'Beta Psi')
        self.assertRedirects(response, reverse('newmember_marks_dashboard'), fetch_redirect_response=False)
        self.assertEqual(self.export(self.nm1, 'Current').status_code, 200)
